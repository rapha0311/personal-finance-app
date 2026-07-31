from openpyxl.chart import PieChart, Reference, BarChart

from io import BytesIO

from openpyxl import Workbook

from openpyxl.styles import numbers

from app.repositories.transaction_repository import (
    get_transactions,
    get_total_income,
    get_total_expenses,
    get_expenses_by_category,
)

from app.repositories.goal_repository import get_goals

from app.repositories.category_repository import get_category_by_id

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.services.transaction_service import get_monthly_financial_report


def generate_excel_report(db):

    workbook = Workbook()

    header_fill = PatternFill("solid", fgColor="1E40AF")

    header_font = Font(color="FFFFFF", bold=True)

    header_alignment = Alignment(horizontal="center")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    income_fill = PatternFill("solid", fgColor="DCFCE7")

    expense_fill = PatternFill("solid", fgColor="FEE2E2")

    income_font = Font(color="166534", bold=True)

    expense_font = Font(color="991B1B", bold=True)

    # ==================================
    # ABA TRANSAÇÕES
    # ==================================

    sheet = workbook.active

    sheet.title = "Transacoes"

    sheet.append(["Data", "Titulo", "Tipo", "Categoria", "Valor"])

    for cell in sheet[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = header_alignment

    transactions = get_transactions(db, skip=0, limit=100000)

    for transaction in transactions:

        category = get_category_by_id(db, transaction.category_id)

        sheet.append(
            [
                str(transaction.transaction_date),
                transaction.title,
                transaction.transaction_type,
                category.name,
                transaction.amount,
            ]
        )

        last_row = sheet.max_row

        sheet[f"E{last_row}"].number_format = "R$ #,##0.00"

        for cell in sheet[last_row]:

            cell.border = thin_border

        if transaction.transaction_type == "income":

            for cell in sheet[last_row]:

                cell.fill = income_fill

            sheet[f"E{last_row}"].font = income_font

        else:

            for cell in sheet[last_row]:

                cell.fill = expense_fill

            sheet[f"E{last_row}"].font = expense_font

    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"

    # ==================================
    # ABA RESUMO
    # ==================================

    summary_sheet = workbook.create_sheet("Resumo")

    summary_sheet.append(["Receitas", "Despesas", "Saldo"])

    for cell in summary_sheet[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = header_alignment

    income = get_total_income(db)

    expenses = get_total_expenses(db)

    balance = income - expenses

    summary_sheet.append([income, expenses, balance])

    for cell in summary_sheet[2]:

        cell.border = thin_border

    summary_sheet["A2"].fill = income_fill
    summary_sheet["A2"].font = income_font

    summary_sheet["B2"].fill = expense_fill
    summary_sheet["B2"].font = expense_font

    if balance >= 0:

        summary_sheet["C2"].fill = income_fill
        summary_sheet["C2"].font = income_font

    else:

        summary_sheet["C2"].fill = expense_fill
        summary_sheet["C2"].font = expense_font

    summary_sheet["A2"].number_format = "R$ #,##0.00"

    summary_sheet["B2"].number_format = "R$ #,##0.00"

    summary_sheet["C2"].number_format = "R$ #,##0.00"

    summary_sheet.auto_filter.ref = summary_sheet.dimensions

    summary_sheet.freeze_panes = "A2"

    # ==================================
    # ABA METAS
    # ==================================

    goals_sheet = workbook.create_sheet("Metas")

    # ... código das metas ...

    goals_sheet.auto_filter.ref = goals_sheet.dimensions

    goals_sheet.freeze_panes = "A2"

    # ==================================
    # ABA INDICADORES
    # ==================================

    indicators_sheet = workbook.create_sheet("Indicadores")

    # Cabeçalho
    indicators_sheet.append(["Categoria", "Total"])

    for cell in indicators_sheet[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = header_alignment

    # Buscar despesas por categoria
    expenses_by_category = get_expenses_by_category(db)

    # Preencher tabela
    for category_name, total in expenses_by_category:

        indicators_sheet.append([category_name, float(total)])

    # Criar gráfico
    pie = PieChart()

    labels = Reference(
        indicators_sheet, min_col=1, min_row=2, max_row=indicators_sheet.max_row
    )

    data = Reference(
        indicators_sheet, min_col=2, min_row=1, max_row=indicators_sheet.max_row
    )

    pie.add_data(data, titles_from_data=True)

    pie.set_categories(labels)

    pie.title = "Despesas por Categoria"

    indicators_sheet.add_chart(pie, "E2")

    monthly_data = get_monthly_financial_report(db)

    start_row = indicators_sheet.max_row + 4

    indicators_sheet.cell(row=start_row, column=1, value="Mês")

    indicators_sheet.cell(row=start_row, column=2, value="Receitas")

    indicators_sheet.cell(row=start_row, column=3, value="Despesas")

    indicators_sheet.cell(row=start_row, column=4, value="Saldo")

    for index, item in enumerate(monthly_data, start=1):

        indicators_sheet.cell(row=start_row + index, column=1, value=item["month"])

        indicators_sheet.cell(row=start_row + index, column=2, value=item["income"])

        indicators_sheet.cell(row=start_row + index, column=3, value=item["expense"])

        indicators_sheet.cell(row=start_row + index, column=4, value=item["balance"])

    chart = BarChart()

    chart.title = "Evolução Financeira Mensal"

    chart.y_axis.title = "Valor"

    chart.x_axis.title = "Mês"

    data = Reference(
        indicators_sheet,
        min_col=2,
        max_col=4,
        min_row=start_row,
        max_row=start_row + len(monthly_data),
    )

    categories = Reference(
        indicators_sheet,
        min_col=1,
        min_row=start_row + 1,
        max_row=start_row + len(monthly_data),
    )

    chart.add_data(data, titles_from_data=True)

    chart.set_categories(categories)

    indicators_sheet.add_chart(chart, "E20")

    # ==================================
    # AJUSTE AUTOMÁTICO DE COLUNAS
    # ==================================

    for worksheet in workbook.worksheets:

        for column_cells in worksheet.columns:

            max_length = 0

            for cell in column_cells:

                try:

                    if cell.value:

                        max_length = max(max_length, len(str(cell.value)))

                except:

                    pass

            adjusted_width = max_length + 5

            column_letter = column_cells[0].column_letter

            worksheet.column_dimensions[column_letter].width = adjusted_width

    # ==================================
    # CENTRALIZAR CONTEÚDO
    # ==================================

    for worksheet in workbook.worksheets:

        for row in worksheet.iter_rows():

            for cell in row:

                cell.alignment = Alignment(horizontal="center")

    file = BytesIO()

    workbook.save(file)

    file.seek(0)

    return file
